import os
import re
import shutil
import json
import openpyxl


class ExcelProcessor:
    def __init__(self):
        self.source_dir = ""
        self.output_dir = ""
    
    def process_project(self, project_name, project_path, excel_path):
        project_data = []
        
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        
        lithology_data = self.read_lithology_sheet(wb)
        
        image_data = self.read_image_sheet(wb)
        
        wb.close()
        
        borehole_image_index = {}
        
        for img_record in image_data:
            borehole = img_record['borehole']
            img_filename = img_record['filename']
            start_depth = img_record['start_depth']
            end_depth = img_record['end_depth']
            
            lithology = self.match_lithology(lithology_data, borehole, start_depth, end_depth)
            
            source_img_path = self.find_image_file(project_path, borehole, img_filename, borehole_image_index)
            
            if source_img_path:
                new_filename = f"{project_name}_{borehole}_{os.path.basename(source_img_path)}"
                dest_img_path = os.path.join(self.output_dir, 'images', new_filename)
                os.makedirs(os.path.dirname(dest_img_path), exist_ok=True)
                
                try:
                    shutil.copy2(source_img_path, dest_img_path)
                    
                    project_data.append({
                        'project': project_name,
                        'borehole': borehole,
                        'image_file': os.path.basename(source_img_path),
                        'new_filename': new_filename,
                        'start_depth': start_depth,
                        'end_depth': end_depth,
                        'lithology': lithology['rock_name'],
                        'lithology_description': lithology['description'],
                        'source_path': source_img_path
                    })
                except Exception as e:
                    print(f"复制图片失败 {source_img_path}: {str(e)}")
        
        return project_data, lithology_data
    
    def read_lithology_sheet(self, wb):
        data = []
        try:
            ws = wb['岩性描述']
            headers = [cell.value for cell in ws[1]]
            
            borehole_idx = headers.index('钻孔编号') if '钻孔编号' in headers else -1
            start_idx = headers.index('起始深度') if '起始深度' in headers else -1
            end_idx = headers.index('终止深度') if '终止深度' in headers else -1
            rock_idx = headers.index('岩石名称') if '岩石名称' in headers else -1
            desc_idx = headers.index('岩性描述') if '岩性描述' in headers else -1
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                    
                data.append({
                    'borehole': row[borehole_idx] if borehole_idx >= 0 else None,
                    'start_depth': float(row[start_idx]) if start_idx >= 0 and row[start_idx] is not None else 0,
                    'end_depth': float(row[end_idx]) if end_idx >= 0 and row[end_idx] is not None else 0,
                    'rock_name': row[rock_idx] if rock_idx >= 0 else '',
                    'description': row[desc_idx] if desc_idx >= 0 else ''
                })
        except KeyError:
            pass
        
        return data
    
    def read_image_sheet(self, wb):
        data = []
        try:
            ws = wb['岩心影像']
            headers = [cell.value for cell in ws[1]]
            
            borehole_idx = headers.index('钻孔编号') if '钻孔编号' in headers else -1
            filename_idx = -1
            for col_name in ['文件名', '图片路径', '图片文件名']:
                if col_name in headers:
                    filename_idx = headers.index(col_name)
                    break
            start_idx = headers.index('起始深度') if '起始深度' in headers else -1
            end_idx = headers.index('终止深度') if '终止深度' in headers else -1
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                
                filename = row[filename_idx] if filename_idx >= 0 else None
                if not filename:
                    continue
                
                filename = str(filename)
                if '//' in filename:
                    filename = filename.split('//', 1)[-1]
                elif '/' in filename:
                    filename = filename.split('/')[-1]
                
                data.append({
                    'borehole': row[borehole_idx] if borehole_idx >= 0 else None,
                    'filename': self.extract_filename(str(filename)),
                    'start_depth': float(row[start_idx]) if start_idx >= 0 and row[start_idx] is not None else 0,
                    'end_depth': float(row[end_idx]) if end_idx >= 0 and row[end_idx] is not None else 0
                })
        except KeyError:
            pass
        
        return data
    
    def extract_filename(self, path):
        if not path:
            return ''
        
        path = path.replace('\\', '/')
        parts = path.split('/')
        filename = parts[-1] if parts else path
        
        filename = re.sub(r'[-_]?\d+[-_]?\d*\.jpg$', '.jpg', filename, flags=re.IGNORECASE)
        filename = re.sub(r'[-_]?\d+[-_]?\d*\.png$', '.png', filename, flags=re.IGNORECASE)
        filename = re.sub(r'[-_]?\d+[-_]?\d*\.bmp$', '.bmp', filename, flags=re.IGNORECASE)
        
        return filename
    
    def find_image_file(self, project_path, borehole, filename, image_index_cache):
        if not borehole or not filename:
            return None
        
        borehole = borehole.strip()
        
        if borehole in image_index_cache:
            index = image_index_cache[borehole]
        else:
            search_dirs = [project_path]
            for d in os.listdir(project_path):
                d_path = os.path.join(project_path, d)
                if os.path.isdir(d_path):
                    search_dirs.append(d_path)
            
            index = {}
            for search_dir in search_dirs:
                if not os.path.isdir(search_dir):
                    continue
                for f in os.listdir(search_dir):
                    if not f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                        continue
                    
                    match = re.match(r'(HC\d+)-(\d+)\.(\w+)', f, re.IGNORECASE)
                    if match:
                        hc_num = match.group(1).upper()
                        suffix = match.group(2)
                        
                        if hc_num not in index:
                            index[hc_num] = {}
                        index[hc_num][suffix] = os.path.join(search_dir, f)
            
            image_index_cache[borehole] = index
        
        for hc, files in index.items():
            for suffix, full_path in files.items():
                if os.path.basename(full_path).upper() == filename.upper():
                    return full_path
        
        return None
    
    def build_image_index(self, borehole_dir):
        index = {}
        
        for f in os.listdir(borehole_dir):
            if not f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                continue
            
            match = re.match(r'(HC\d+)-(\d+)\.(\w+)', f, re.IGNORECASE)
            if match:
                hc_num = match.group(1).upper()
                suffix = match.group(2)
                
                if hc_num not in index:
                    index[hc_num] = {}
                index[hc_num][suffix] = os.path.join(borehole_dir, f)
        
        return index
    
    def match_lithology(self, lithology_data, borehole, start_depth, end_depth):
        for lith in lithology_data:
            if lith['borehole'] != borehole:
                continue
            
            lith_start = lith['start_depth']
            lith_end = lith['end_depth']
            
            if start_depth >= lith_start and end_depth <= lith_end:
                return lith
            
            if start_depth < lith_end and end_depth > lith_start:
                return lith
        
        return {
            'borehole': borehole,
            'rock_name': '',
            'description': '',
            'start_depth': 0,
            'end_depth': 0
        }
