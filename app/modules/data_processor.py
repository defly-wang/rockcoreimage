import os
import re
import shutil
import json
from pathlib import Path
import openpyxl
from PyQt6.QtCore import QObject, pyqtSignal


class DataProcessor(QObject):
    progress_updated = pyqtSignal(int, str)
    processing_finished = pyqtSignal(str, dict)
    error_occurred = pyqtSignal(str)
    
    def __init__(self):
        super().__init__()
        self.source_dir = ""
        self.output_dir = ""
        
    def classify_lithology(self, json_file, config_file, output_file):
        with open(config_file, 'r', encoding='utf-8') as f:
            rock_types = json.load(f)
        
        rocks = set()
        for rock in rock_types['rocks']:
            rocks.add(rock['name'])
            rocks.update(rock.get('aliases', []))
        rocks = sorted(rocks, key=lambda x: -len(x))
        
        self.progress_updated.emit(0, "正在加载数据...")
        
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        total = len(data)
        
        def extract_last_keyword(lithology):
            lithology = lithology.strip()
            lithology = re.sub(r'[（(][^）)]*[）)]$', '', lithology).strip()
            for rock in rocks:
                if lithology.endswith(rock):
                    return rock
            return ''
        
        mapping = {}
        unmatched = {}
        for idx, item in enumerate(data):
            if idx % 100 == 0:
                self.progress_updated.emit(int(idx / total * 100), f"正在分类... {idx}/{total}")
            
            lithology = item.get('lithology', '')
            keyword = extract_last_keyword(lithology)
            item['岩性名称'] = keyword
            
            if keyword:
                if keyword not in mapping:
                    mapping[keyword] = {'lithologies': set(), 'count': 0}
                mapping[keyword]['lithologies'].add(lithology)
                mapping[keyword]['count'] += 1
            else:
                if lithology not in unmatched:
                    unmatched[lithology] = 0
                unmatched[lithology] += 1
        
        self.progress_updated.emit(95, "正在保存文件...")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        self.progress_updated.emit(100, "分类完成")
        
        matched = sum(v['count'] for v in mapping.values())
        
        stats = {
            'total': total,
            'matched': matched,
            'types': len(mapping),
            'mapping': mapping,
            'unmatched': unmatched
        }
        
        self.processing_finished.emit(output_file, stats)
        
    def process(self, source_dir, output_dir):
        self.source_dir = source_dir
        self.output_dir = output_dir
        
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(os.path.join(output_dir, 'images'), exist_ok=True)
        
        self.progress_updated.emit(0, "正在扫描项目目录...")
        
        project_dirs = [d for d in os.listdir(source_dir) 
                      if os.path.isdir(os.path.join(source_dir, d)) and not d.startswith('.')]
        
        total_projects = len(project_dirs)
        if total_projects == 0:
            self.progress_updated.emit(100, "未找到项目")
            self.processing_finished.emit("", {'total_images': 0, 'total_projects': 0})
            return
        
        self.progress_updated.emit(5, f"共发现 {total_projects} 个项目")
        
        all_data = []
        
        for idx, project in enumerate(sorted(project_dirs)):
            progress_base = int(10 + (idx / total_projects) * 80)
            self.progress_updated.emit(
                progress_base,
                f"处理项目 [{idx+1}/{total_projects}]: {project}"
            )
            
            project_path = os.path.join(source_dir, project)
            
            self.progress_updated.emit(
                progress_base,
                f"正在读取项目 {project} 的Excel文件..."
            )
            
            excel_files = [f for f in os.listdir(project_path) 
                         if f.endswith('.xlsx') and not f.startswith('~') and not f.startswith('.~')]
            
            if not excel_files:
                self.progress_updated.emit(
                    progress_base,
                    f"项目 {project} 无Excel文件，跳过"
                )
                continue
            
            excel_path = os.path.join(project_path, excel_files[0])
            
            try:
                self.progress_updated.emit(
                    progress_base + 2,
                    f"正在解析 {excel_files[0]}..."
                )
                project_data = self.process_project(project, project_path, excel_path)
                all_data.extend(project_data)
                self.progress_updated.emit(
                    progress_base + 5,
                    f"项目 {project} 处理完成，获取 {len(project_data)} 条记录"
                )
            except Exception as e:
                self.error_occurred.emit(f"处理项目 {project} 时出错: {str(e)}")
        
        self.progress_updated.emit(95, "正在保存数据文件...")
        
        lithology_stats = {}
        for item in all_data:
            lith = item.get('lithology', '')
            if lith:
                if lith not in lithology_stats:
                    lithology_stats[lith] = {'count': 0, 'description': item.get('lithology_description', '')}
                lithology_stats[lith]['count'] += 1
        
        output_file = os.path.join(output_dir, 'image_descriptions.json')
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(all_data, f, ensure_ascii=False, indent=2)
        
        self.progress_updated.emit(100, "处理完成")
        self.processing_finished.emit(output_file, {
            'total_images': len(all_data),
            'total_projects': total_projects,
            'lithology_stats': lithology_stats
        })
    
    def process_project(self, project_name, project_path, excel_path):
        project_data = []
        
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        
        lithology_data = self.read_lithology_sheet(wb)
        
        image_data = self.read_image_sheet(wb)
        
        wb.close()
        
        borehole_image_index = {}
        
        for img_record in image_data:
            borehole = img_record['borehole']
            img_filename = img_record['filename']
            start_depth = img_record['start_depth']
            end_depth = img_record['end_depth']
            
            lithology = self.match_lithology(lithology_data, borehole, start_depth, end_depth)
            
            source_img_path = self.find_image_file(project_path, borehole, img_filename, borehole_image_index)
            
            if source_img_path:
                new_filename = f"{project_name}_{borehole}_{os.path.basename(source_img_path)}"
                dest_img_path = os.path.join(self.output_dir, 'images', new_filename)
                os.makedirs(os.path.dirname(dest_img_path), exist_ok=True)
                
                try:
                    shutil.copy2(source_img_path, dest_img_path)
                    
                    project_data.append({
                        'project': project_name,
                        'borehole': borehole,
                        'image_file': os.path.basename(source_img_path),
                        'new_filename': new_filename,
                        'start_depth': start_depth,
                        'end_depth': end_depth,
                        'lithology': lithology['rock_name'],
                        'lithology_description': lithology['description'],
                        'source_path': source_img_path
                    })
                except Exception as e:
                    self.error_occurred.emit(f"复制图片失败 {source_img_path}: {str(e)}")
        
        return project_data
    
    def read_lithology_sheet(self, wb):
        data = []
        try:
            ws = wb['岩性描述']
            headers = [cell.value for cell in ws[1]]
            
            borehole_idx = headers.index('钻孔编号') if '钻孔编号' in headers else -1
            start_idx = headers.index('起始深度') if '起始深度' in headers else -1
            end_idx = headers.index('终止深度') if '终止深度' in headers else -1
            rock_idx = headers.index('岩石名称') if '岩石名称' in headers else -1
            desc_idx = headers.index('岩性描述') if '岩性描述' in headers else -1
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                    
                data.append({
                    'borehole': row[borehole_idx] if borehole_idx >= 0 else None,
                    'start_depth': float(row[start_idx]) if start_idx >= 0 and row[start_idx] is not None else 0,
                    'end_depth': float(row[end_idx]) if end_idx >= 0 and row[end_idx] is not None else 0,
                    'rock_name': row[rock_idx] if rock_idx >= 0 else '',
                    'description': row[desc_idx] if desc_idx >= 0 else ''
                })
        except KeyError:
            pass
        
        return data
    
    def read_image_sheet(self, wb):
        data = []
        try:
            ws = wb['岩心影像']
            headers = [cell.value for cell in ws[1]]
            
            borehole_idx = headers.index('钻孔编号') if '钻孔编号' in headers else -1
            start_idx = headers.index('起始深度') if '起始深度' in headers else -1
            end_idx = headers.index('终止深度') if '终止深度' in headers else -1
            path_idx = headers.index('图片路径') if '图片路径' in headers else -1
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                
                path = row[path_idx] if path_idx >= 0 else ''
                filename = self.extract_filename(path)
                
                data.append({
                    'borehole': row[borehole_idx] if borehole_idx >= 0 else None,
                    'start_depth': float(row[start_idx]) if start_idx >= 0 and row[start_idx] is not None else 0,
                    'end_depth': float(row[end_idx]) if end_idx >= 0 and row[end_idx] is not None else 0,
                    'filename': filename,
                    'path': path
                })
        except KeyError:
            pass
        
        return data
    
    def extract_filename(self, path):
        if not path:
            return ''
        
        path = path.replace('\\', '/')
        if '//' in path:
            path = path.split('//', 1)[1]
        
        return os.path.basename(path)
    
    def find_image_file(self, project_path, borehole, filename, image_index_cache):
        borehole_dir = os.path.join(project_path, borehole)
        
        if not os.path.exists(borehole_dir):
            return None
        
        if borehole not in image_index_cache:
            image_index_cache[borehole] = self.build_image_index(borehole_dir)
        
        index = image_index_cache[borehole]
        
        match = re.match(r'(HC\d+)-(\d+)\.(\w+)', filename, re.IGNORECASE)
        if match:
            hc_num = match.group(1).upper()
            suffix = match.group(2)
            ext = match.group(3).lower()
            
            if hc_num in index and suffix in index[hc_num]:
                return index[hc_num][suffix]
        
        for hc, files in index.items():
            for suffix, full_path in files.items():
                if os.path.basename(full_path).upper() == filename.upper():
                    return full_path
        
        return None
    
    def build_image_index(self, borehole_dir):
        index = {}
        
        for f in os.listdir(borehole_dir):
            if not f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                continue
            
            match = re.match(r'(HC\d+)-(\d+)\.(\w+)', f, re.IGNORECASE)
            if match:
                hc_num = match.group(1).upper()
                suffix = match.group(2)
                
                if hc_num not in index:
                    index[hc_num] = {}
                index[hc_num][suffix] = os.path.join(borehole_dir, f)
        
        return index
    
    def match_lithology(self, lithology_data, borehole, start_depth, end_depth):
        for lith in lithology_data:
            if lith['borehole'] != borehole:
                continue
            
            lith_start = lith['start_depth']
            lith_end = lith['end_depth']
            
            if start_depth >= lith_start and end_depth <= lith_end:
                return lith
            
            if start_depth < lith_end and end_depth > lith_start:
                return lith
        
        return {
            'borehole': borehole,
            'rock_name': '',
            'description': '',
            'start_depth': 0,
            'end_depth': 0
        }
    
    def analyze_alteration(self, json_file, alteration_config_file, output_file):
        with open(alteration_config_file, 'r', encoding='utf-8') as f:
            alteration_config = json.load(f)
        
        alterations = alteration_config.get('alterations', [])
        
        alteration_keywords = {}
        for alt in alterations:
            alt_name = alt['name']
            keywords = [alt_name]
            keywords.extend(alt.get('aliases', []))
            alteration_keywords[alt_name] = {
                'keywords': set(keywords),
                'category': alt.get('category', ''),
                'importance': alt.get('importance', False)
            }
        
        alteration_keywords = dict(sorted(alteration_keywords.items(), key=lambda x: -len(x[1]['keywords'])))
        
        self.progress_updated.emit(0, "正在加载数据...")
        
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        total = len(data)
        
        def detect_alterations(lithology):
            if not lithology:
                return []
            detected = []
            lithology_lower = lithology.lower()
            for alt_name, alt_info in alteration_keywords.items():
                for keyword in alt_info['keywords']:
                    if keyword and keyword.lower() in lithology_lower:
                        detected.append(alt_name)
                        break
            return detected
        
        alteration_stats = {}
        total_alterations = 0
        
        for idx, item in enumerate(data):
            if idx % 100 == 0:
                self.progress_updated.emit(int(idx / total * 80), f"正在分析蚀变... {idx}/{total}")
            
            lithology = item.get('lithology', '')
            
            detected = detect_alterations(lithology)
            item['蚀变类型'] = ", ".join(detected) if detected else ""
            
            for alt_name in detected:
                if alt_name not in alteration_stats:
                    alteration_stats[alt_name] = {
                        'category': alteration_keywords[alt_name]['category'],
                        'importance': alteration_keywords[alt_name]['importance'],
                        'count': 0
                    }
                alteration_stats[alt_name]['count'] += 1
                total_alterations += 1
        
        self.progress_updated.emit(90, "正在统计结果...")
        
        importance_stats = {
            'important': {},
            'other': {}
        }
        category_stats = {}
        
        for alt_name, info in alteration_stats.items():
            target = importance_stats['important'] if info['importance'] else importance_stats['other']
            target[alt_name] = info['count']
            
            if info['category'] not in category_stats:
                category_stats[info['category']] = 0
            category_stats[info['category']] += info['count']
        
        self.progress_updated.emit(95, "正在保存文件...")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        self.progress_updated.emit(100, "蚀变分析完成")
        
        stats = {
            'total_records': total,
            'records_with_alteration': sum(1 for item in data if item.get('蚀变类型')),
            'total_alterations': total_alterations,
            'alteration_types': len(alteration_stats),
            'important_alterations': importance_stats['important'],
            'other_alterations': importance_stats['other'],
            'category_stats': category_stats
        }
        
        self.processing_finished.emit(output_file, stats)
    
    def process_html_project(self, source_dir, output_dir):
        self.source_dir = source_dir
        self.output_dir = output_dir
        
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(os.path.join(output_dir, 'images'), exist_ok=True)
        
        self.progress_updated.emit(0, "正在扫描项目目录...")
        
        project_dirs = [d for d in os.listdir(source_dir) 
                      if os.path.isdir(os.path.join(source_dir, d)) and not d.startswith('.')]
        
        total_projects = len(project_dirs)
        if total_projects == 0:
            self.progress_updated.emit(100, "未找到项目")
            self.processing_finished.emit("", {'total_images': 0, 'total_projects': 0})
            return
        
        self.progress_updated.emit(5, f"共发现 {total_projects} 个项目")
        
        all_data = []
        
        for idx, project in enumerate(sorted(project_dirs)):
            progress_base = int(10 + (idx / total_projects) * 80)
            self.progress_updated.emit(
                progress_base,
                f"处理项目 [{idx+1}/{total_projects}]: {project}"
            )
            
            project_path = os.path.join(source_dir, project)
            
            offline_dir = os.path.join(project_path, '离线成果展示')
            borehole_dir = project_path
            
            found = False
            for d in os.listdir(project_path):
                d_path = os.path.join(project_path, d)
                if os.path.isdir(d_path) and d.endswith('筒次'):
                    borehole_dir = project_path
                    found = True
                    break
                if os.path.isdir(d_path):
                    for sub_d in os.listdir(d_path):
                        if sub_d.endswith('筒次'):
                            borehole_dir = d_path
                            found = True
                            break
                if found:
                    break
            
            if not os.path.exists(offline_dir):
                self.progress_updated.emit(progress_base, f"项目 {project} 无离线成果展示目录，跳过")
                continue
            
            white_light_html = os.path.join(offline_dir, '白光平扫相册.html')
            histogram_html = os.path.join(offline_dir, '综合柱状图.html')
            
            if not os.path.exists(white_light_html) or not os.path.exists(histogram_html):
                self.progress_updated.emit(progress_base, f"项目 {project} 缺少HTML文件，跳过")
                continue
            
            try:
                self.progress_updated.emit(progress_base + 2, f"正在解析 {project} 的HTML文件...")
                project_data = self.process_html_files(project, borehole_dir, white_light_html, histogram_html)
                all_data.extend(project_data)
                self.progress_updated.emit(
                    progress_base + 5,
                    f"项目 {project} 处理完成，获取 {len(project_data)} 条记录"
                )
            except Exception as e:
                self.error_occurred.emit(f"处理项目 {project} 时出错: {str(e)}")
        
        self.progress_updated.emit(95, "正在保存数据文件...")
        
        lithology_stats = {}
        for item in all_data:
            lith = item.get('lithology', '')
            if lith:
                if lith not in lithology_stats:
                    lithology_stats[lith] = {'count': 0, 'description': item.get('lithology_description', '')}
                lithology_stats[lith]['count'] += 1
        
        output_file = os.path.join(output_dir, 'image_descriptions.json')
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(all_data, f, ensure_ascii=False, indent=2)
        
        self.progress_updated.emit(100, "处理完成")
        self.processing_finished.emit(output_file, {
            'total_images': len(all_data),
            'total_projects': total_projects,
            'lithology_stats': lithology_stats
        })
    
    def process_html_files(self, project_name, borehole_dir, white_light_html, histogram_html):
        project_data = []
        
        image_info = self.parse_white_light_html(white_light_html)
        
        lithology_info = self.parse_histogram_html(histogram_html)
        
        lithology_map = {}
        for lith in lithology_info:
            key = (lith['start_depth'], lith['end_depth'])
            lithology_map[key] = lith
        
        for img in image_info:
            full_path = img['path']
            full_path = full_path.replace('\\', '/')
            if '//' in full_path:
                full_path = full_path.split('//', 1)[1]
            
            filename = os.path.basename(full_path)
            start_depth = img['start_depth']
            end_depth = img['end_depth']
            
            matched_lith = None
            for (s, e), lith in lithology_map.items():
                if start_depth >= s and end_depth <= e:
                    matched_lith = lith
                    break
                if start_depth < e and end_depth > s:
                    matched_lith = lith
                    break
            
            source_img_path = self.find_image_in_tongci(borehole_dir, filename)
        
        image_info = self.parse_white_light_html(white_light_html)
        
        lithology_info = self.parse_histogram_html(histogram_html)
        
        rocks = self.load_rock_types()
        
        lithology_map = {}
        for lith in lithology_info:
            key = (lith['start_depth'], lith['end_depth'])
            lithology_map[key] = lith
        
        for img in image_info:
            full_path = img['path']
            full_path = full_path.replace('\\', '/')
            if '//' in full_path:
                full_path = full_path.split('//', 1)[1]
            
            filename = os.path.basename(full_path)
            start_depth = img['start_depth']
            end_depth = img['end_depth']
            
            matched_lith = None
            for (s, e), lith in lithology_map.items():
                if start_depth >= s and end_depth <= e:
                    matched_lith = lith
                    break
                if start_depth < e and end_depth > s:
                    matched_lith = lith
                    break
            
            lithology = matched_lith['rock_name'] if matched_lith else ''
            lithology_description = matched_lith['description'] if matched_lith else ''
            
            if lithology_description:
                refined_lithology = self.analyze_lithology_from_description(
                    lithology, lithology_description, start_depth, end_depth, rocks
                )
                if refined_lithology:
                    lithology = refined_lithology
            
            source_img_path = self.find_image_in_tongci(borehole_dir, filename)
            
            if source_img_path:
                new_filename = f"{project_name}_{os.path.basename(source_img_path)}"
                dest_img_path = os.path.join(self.output_dir, 'images', new_filename)
                os.makedirs(os.path.dirname(dest_img_path), exist_ok=True)
                
                try:
                    shutil.copy2(source_img_path, dest_img_path)
                    
                    project_data.append({
                        'project': project_name,
                        'borehole': project_name,
                        'image_file': os.path.basename(source_img_path),
                        'new_filename': new_filename,
                        'start_depth': start_depth,
                        'end_depth': end_depth,
                        'lithology': lithology,
                        'lithology_description': lithology_description,
                        'source_path': source_img_path
                    })
                except Exception as e:
                    self.error_occurred.emit(f"复制图片失败 {source_img_path}: {str(e)}")
        
        return project_data
    
    def load_rock_types(self):
        config_files = [
            'config/rock_types.json',
            'config/rock_types_comprehensive.json',
            'config/rock_types_flat.json'
        ]
        
        rocks = set()
        for config_file in config_files:
            try:
                with open(config_file, 'r', encoding='utf-8') as f:
                    config = json.load(f)
                
                if 'rocks' in config:
                    for rock in config['rocks']:
                        rocks.add(rock['name'])
                        rocks.update(rock.get('aliases', []))
                elif 'categories' in config:
                    for category, cat_info in config['categories'].items():
                        for rock_list in [cat_info.get('important', []), cat_info.get('general', [])]:
                            rocks.update(rock_list)
            except:
                pass
        
        return sorted(rocks, key=lambda x: -len(x))
    
    def analyze_lithology_from_description(self, current_lithology, description, start_depth, end_depth, rocks):
        if not description:
            return current_lithology
        
        all_depth_ranges = []
        
        wei_patterns = [
            r'为([\u4e00-\u9fa5]+岩)',
            r'为([\u4e00-\u9fa5]+)',
            r'是([\u4e00-\u9fa5]+岩)',
            r'是([\u4e00-\u9fa5]+)',
        ]
        
        for wei_pattern in wei_patterns:
            for wei_match in re.finditer(wei_pattern, description):
                lithology = wei_match.group(1)
                
                lithology_match = None
                for rock in rocks:
                    if rock == lithology:
                        lithology_match = rock
                        break
                    if rock in lithology or lithology in rock:
                        lithology_match = rock
                        break
                
                if lithology_match is None:
                    if any(rock in lithology or lithology in rock for rock in rocks):
                        for rock in rocks:
                            if rock in lithology or lithology in rock:
                                lithology_match = rock
                                break
                    else:
                        lithology_match = lithology
                
                if lithology_match:
                    text_before = description[:wei_match.start()]
                    
                    sentences = re.split(r'[。；]', text_before)
                    
                    for sentence in reversed(sentences):
                        if not sentence.strip():
                            continue
                        
                        depth_pattern = r'([\d.]+)[m米]?[-–—]([\d.]+)[m米]?'
                        all_depths = list(re.finditer(depth_pattern, sentence))
                        
                        for dm in all_depths:
                            all_depth_ranges.append((
                                lithology_match,
                                dm.group(1),
                                dm.group(2)
                            ))
                        
                        if all_depths:
                            break
        
        for rock, desc_start, desc_end_str in all_depth_ranges:
            try:
                desc_start = float(desc_start)
                desc_end = float(desc_end_str)
                
                overlap_start = max(start_depth, desc_start)
                overlap_end = min(end_depth, desc_end)
                overlap = max(0, overlap_end - overlap_start)
                img_range = end_depth - start_depth
                
                if img_range > 0 and overlap / img_range >= 0.6:
                    return rock
            except:
                pass
        
        return current_lithology
    
    def parse_white_light_html(self, html_file):
        import re
        
        info_js_path = os.path.join(os.path.dirname(html_file), '白光平扫相册_files', 'Info.js')
        if not os.path.exists(info_js_path):
            return []
        
        with open(info_js_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        match = re.search(r'var\s+imgjson\s*=\s*(\[\[.*?\]\]);', content, re.DOTALL)
        if not match:
            return []
        
        json_str = match.group(1)
        
        json_str = json_str.replace('\\', '\\\\')
        
        try:
            data = json.loads(json_str)
        except json.JSONDecodeError:
            data = []
        
        result = []
        for batch in data:
            for item in batch:
                yxbh = item.get('Yxbh', '')
                qsjs = float(item.get('Qsjs', 0))
                zzjs = float(item.get('Zzjs', 0))
                txlj = item.get('Txlj', '')
                
                filename = os.path.basename(txlj)
                
                result.append({
                    'filename': filename,
                    'yxbh': yxbh,
                    'start_depth': qsjs,
                    'end_depth': zzjs,
                    'path': txlj
                })
        
        return result
    
    def parse_histogram_html(self, html_file):
        sys_js_path = os.path.join(os.path.dirname(html_file), '综合柱状图_files', 'SysHistogramInfo.js')
        if not os.path.exists(sys_js_path):
            return []
        
        with open(sys_js_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        import re
        match = re.search(r'var\s+json\s*=\s*(\[.*?\]);', content, re.DOTALL)
        if not match:
            return []
        
        json_str = match.group(1)
        json_str = json_str.replace('\\', '\\\\')
        
        try:
            data = json.loads(json_str)
        except json.JSONDecodeError:
            return []
        
        if len(data) < 2:
            return []
        
        lithology_batch = data[1]
        
        result = []
        for item in lithology_batch:
            if item.get('Type') != '岩心描述':
                continue
                
            qsjs = item.get('QSJD', 0)
            zzjs = item.get('ZZJD', 0)
            rock_name = item.get('YSMC', '')
            description = item.get('MS', '')
            
            if description:
                description = description.replace('||', '')
                description = description.replace('\\n', '\n')
                description = description.replace('\\r', '\r')
                description = description.replace('\\t', '\t')
                description = description.strip()
            
            result.append({
                'start_depth': float(qsjs) if qsjs else 0,
                'end_depth': float(zzjs) if zzjs else 0,
                'rock_name': rock_name,
                'description': description
            })
        
        return result
    
    def find_image_in_tongci(self, borehole_dir, filename):
        if not borehole_dir or not os.path.exists(borehole_dir):
            return None
        
        import re
        def extract_number(d):
            match = re.match(r'第(\d+)筒次', d)
            return int(match.group(1)) if match else 0
        
        tongci_dirs = [d for d in os.listdir(borehole_dir) 
                      if d.endswith('筒次') and os.path.isdir(os.path.join(borehole_dir, d))]
        
        for tongci in sorted(tongci_dirs, key=extract_number):
            tongci_path = os.path.join(borehole_dir, tongci)
            original_img_dir = os.path.join(tongci_path, '原始图像')
            
            if not os.path.exists(original_img_dir):
                continue
            
            for root, dirs, files in os.walk(original_img_dir):
                for f in files:
                    if f.lower() == filename.lower():
                        return os.path.join(root, f)
        
        return None
