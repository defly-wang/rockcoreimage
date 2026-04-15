import os
from PyQt6.QtWidgets import QMessageBox, QTableWidgetItem


class ProcessHandler:
    def __init__(self, main_window):
        self.main_window = main_window
    
    def select_source_directory(self):
        from PyQt6.QtWidgets import QFileDialog
        folder = QFileDialog.getExistingDirectory(self.main_window, "选择数据目录")
        if folder:
            self.main_window.source_directory = folder
            self.main_window.source_path_label.setText(os.path.basename(folder))
            self.main_window.process_log.append(f"已选择数据源: {folder}")
    
    def select_output_directory(self):
        from PyQt6.QtWidgets import QFileDialog
        folder = QFileDialog.getExistingDirectory(self.main_window, "选择输出目录")
        if folder:
            self.main_window.output_directory = folder
            self.main_window.output_path_label.setText(os.path.basename(folder))
            self.main_window.process_log.append(f"已选择输出目录: {folder}")
    
    def show_process_stats(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择处理结果文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        import json
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            QMessageBox.warning(self.main_window, "错误", f"无法读取文件: {str(e)}")
            return
        
        self.main_window.process_log.clear()
        self.main_window.process_log.append(f"正在加载: {json_file}")
        
        lithology_stats = {}
        lith_order = {}
        proj_order = {}
        for idx, item in enumerate(data):
            lith = item.get('lithology', '')
            proj = item.get('project', '')
            if lith:
                if proj not in proj_order:
                    proj_order[proj] = len(proj_order)
                key = (lith, proj)
                if key not in lithology_stats:
                    lithology_stats[key] = {'lithology': lith, 'project': proj, 'count': 0, 'description': item.get('lithology_description', '')}
                    lith_order[key] = len(lith_order)
                lithology_stats[key]['count'] += 1
        
        for key in lithology_stats:
            lithology_stats[key]['order'] = lith_order[key]
            lithology_stats[key]['proj_order'] = proj_order[key[1]]
        
        self.main_window.process_table.setColumnCount(4)
        self.main_window.process_table.setHorizontalHeaderLabels(["项目", "岩性名称", "图片数", "岩性描述"])
        self.main_window.process_table.setColumnWidth(0, 120)
        self.main_window.process_table.setColumnWidth(1, 120)
        self.main_window.process_table.setColumnWidth(2, 80)
        self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
        
        sorted_lith = sorted(lithology_stats.values(), key=lambda x: (x.get('proj_order', 0), x.get('order', 0)))
        self.main_window.process_table.setRowCount(len(sorted_lith))
        
        for i, info in enumerate(sorted_lith):
            self.main_window.process_table.setItem(i, 0, QTableWidgetItem(info['project']))
            self.main_window.process_table.setItem(i, 1, QTableWidgetItem(info['lithology']))
            self.main_window.process_table.setItem(i, 2, QTableWidgetItem(str(info['count'])))
            desc = info['description']
            self.main_window.process_table.setItem(i, 3, QTableWidgetItem(desc.replace('\n', ' ') if desc else ''))
        
        self.main_window.process_table.resizeRowsToContents()
        self.main_window.process_log.append(f"已加载 {len(data)} 条记录，按项目/岩性统计共 {len(sorted_lith)} 项")
    
    def start_data_processing(self):
        if not hasattr(self.main_window, 'source_directory') or not self.main_window.source_directory:
            QMessageBox.warning(self.main_window, "警告", "请先选择数据源目录")
            return
        
        if not hasattr(self.main_window, 'output_directory') or not self.main_window.output_directory:
            QMessageBox.warning(self.main_window, "警告", "请先选择输出目录")
            return
        
        process_type = getattr(self.main_window, 'process_type', 'excel')
        
        self.main_window.process_btn.setEnabled(False)
        if hasattr(self.main_window, 'classify_btn'):
            self.main_window.classify_btn.setEnabled(False)
        if hasattr(self.main_window, 'alteration_btn'):
            self.main_window.alteration_btn.setEnabled(False)
        if hasattr(self.main_window, 'stats_btn'):
            self.main_window.stats_btn.setEnabled(False)
        self.main_window.process_status_label.setText("正在初始化...")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #FF9800;
            padding: 8px;
            background-color: #FFF3E0;
            border: 1px solid #FF9800;
            border-radius: 4px;
        """)
        self.main_window.process_progress.setValue(0)
        self.main_window.process_log.clear()
        
        if process_type == 'html':
            self.main_window.process_log.append("开始处理HTML数据...")
        else:
            self.main_window.process_log.append("开始处理数据...")
        
        try:
            self.main_window.data_processor.progress_updated.disconnect()
            self.main_window.data_processor.processing_finished.disconnect()
            self.main_window.data_processor.error_occurred.disconnect()
        except TypeError:
            pass
        
        self.main_window.data_processor.progress_updated.connect(self.on_processing_progress)
        self.main_window.data_processor.processing_finished.connect(self.on_processing_finished)
        self.main_window.data_processor.error_occurred.connect(self.on_processing_error)
        
        from threading import Thread
        process_type = getattr(self.main_window, 'process_type', 'excel')
        
        if process_type == 'html':
            self.main_window.process_thread = Thread(
                target=self.main_window.data_processor.process_html_project,
                args=(self.main_window.source_directory, self.main_window.output_directory),
                daemon=True
            )
        else:
            self.main_window.process_thread = Thread(
                target=self.main_window.data_processor.process,
                args=(self.main_window.source_directory, self.main_window.output_directory),
                daemon=True
            )
        self.main_window.process_thread.start()
    
    def on_processing_progress(self, value, message):
        self.main_window.process_progress.setValue(value)
        self.main_window.process_log.append(message)
        self.main_window.status_bar.showMessage(message)
        
        if value < 30:
            self.main_window.process_status_label.setText("正在扫描项目...")
            self.main_window.process_status_label.setStyleSheet("""
                font-size: 14px;
                font-weight: bold;
                color: #2196F3;
                padding: 8px;
                background-color: #E3F2FD;
                border: 1px solid #2196F3;
                border-radius: 4px;
            """)
        elif value < 90:
            self.main_window.process_status_label.setText("正在处理数据...")
            self.main_window.process_status_label.setStyleSheet("""
                font-size: 14px;
                font-weight: bold;
                color: #FF9800;
                padding: 8px;
                background-color: #FFF3E0;
                border: 1px solid #FF9800;
                border-radius: 4px;
            """)
        else:
            self.main_window.process_status_label.setText("正在保存结果...")
            self.main_window.process_status_label.setStyleSheet("""
                font-size: 14px;
                font-weight: bold;
                color: #9C27B0;
                padding: 8px;
                background-color: #F3E5F5;
                border: 1px solid #9C27B0;
                border-radius: 4px;
            """)
    
    def on_processing_finished(self, output_file, stats):
        self.main_window.process_progress.setValue(100)
        
        if stats.get('total_images', 0) > 0:
            self.main_window.process_status_label.setText(f"处理完成 - 共 {stats['total_images']} 张图片，{stats['total_projects']} 个项目")
            self.main_window.process_status_label.setStyleSheet("""
                font-size: 14px;
                font-weight: bold;
                color: #4CAF50;
                padding: 8px;
                background-color: #E8F5E9;
                border: 1px solid #4CAF50;
                border-radius: 4px;
            """)
            self.main_window.process_btn.setEnabled(True)
            self.main_window.classify_btn.setEnabled(True)
            if hasattr(self.main_window, 'alteration_btn'):
                self.main_window.alteration_btn.setEnabled(True)
            if hasattr(self.main_window, 'stats_btn'):
                self.main_window.stats_btn.setEnabled(True)
            self.main_window.classify_output_file = output_file
            
            lithology_stats = stats.get('lithology_stats', {})
            
            self.main_window.process_table.setColumnCount(4)
            self.main_window.process_table.setHorizontalHeaderLabels(["项目", "岩性名称", "图片数", "岩性描述"])
            self.main_window.process_table.setColumnWidth(0, 120)
            self.main_window.process_table.setColumnWidth(1, 120)
            self.main_window.process_table.setColumnWidth(2, 80)
            self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
            
            sorted_lith = sorted(lithology_stats.values(), key=lambda x: (x.get('proj_order', 0), x.get('order', 0)))
            self.main_window.process_table.setRowCount(len(sorted_lith))
            
            for i, info in enumerate(sorted_lith):
                self.main_window.process_table.setItem(i, 0, QTableWidgetItem(info['project']))
                self.main_window.process_table.setItem(i, 1, QTableWidgetItem(info['lithology']))
                self.main_window.process_table.setItem(i, 2, QTableWidgetItem(str(info['count'])))
                desc = info['description']
                self.main_window.process_table.setItem(i, 3, QTableWidgetItem(desc.replace('\n', ' ')))
            
            self.main_window.process_table.resizeRowsToContents()
        else:
            self.main_window.process_status_label.setText("处理完成 - 未找到数据")
            self.main_window.process_status_label.setStyleSheet("""
                font-size: 14px;
                font-weight: bold;
                color: #9E9E9E;
                padding: 8px;
                background-color: #F5F5F5;
                border: 1px solid #9E9E9E;
                border-radius: 4px;
            """)
            self.main_window.process_btn.setEnabled(True)
            self.main_window.classify_btn.setEnabled(True)
        
        self.main_window.process_log.append(f"处理完成!")
        self.main_window.process_log.append(f"共处理图片: {stats.get('total_images', 0)} 张")
        self.main_window.process_log.append(f"项目数量: {stats.get('total_projects', 0)} 个")
        self.main_window.process_log.append(f"描述文件: {output_file}")
        
        self.main_window.status_bar.showMessage("处理完成")
        
        if stats.get('total_images', 0) > 0:
            QMessageBox.information(self.main_window, "处理完成", 
                f"共处理图片: {stats['total_images']} 张\n"
                f"项目数量: {stats['total_projects']} 个\n"
                f"描述文件: {output_file}")
    
    def on_processing_error(self, error_msg):
        self.main_window.process_log.append(f"错误: {error_msg}")
        self.main_window.process_btn.setEnabled(True)
        self.main_window.classify_btn.setEnabled(True)
    
    def start_lithology_classify(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择JSON文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        app_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        project_root = os.path.dirname(app_dir)
        config_file = os.path.join(project_root, 'config', 'rock_types_flat.json')
        
        if not os.path.exists(config_file):
            QMessageBox.warning(self.main_window, "错误", f"找不到配置文件: {config_file}\n请检查config目录是否存在")
            return
        
        output_file = json_file.replace('.json', '_classified.json')
        
        self.main_window.process_btn.setEnabled(False)
        self.main_window.classify_btn.setEnabled(False)
        self.main_window.process_status_label.setText("正在初始化...")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #FF9800;
            padding: 8px;
            background-color: #FFF3E0;
            border: 1px solid #FF9800;
            border-radius: 4px;
        """)
        self.main_window.process_progress.setValue(0)
        self.main_window.process_log.clear()
        self.main_window.process_log.append("开始岩性分类...")
        
        try:
            self.main_window.data_processor.progress_updated.disconnect()
            self.main_window.data_processor.processing_finished.disconnect()
            self.main_window.data_processor.error_occurred.disconnect()
        except TypeError:
            pass
        
        self.main_window.data_processor.progress_updated.connect(self.on_classify_progress)
        self.main_window.data_processor.processing_finished.connect(self.on_classify_finished)
        self.main_window.data_processor.error_occurred.connect(self.on_processing_error)
        
        from threading import Thread
        self.main_window.classify_thread = Thread(
            target=self.main_window.data_processor.classify_lithology,
            args=(json_file, config_file, output_file),
            daemon=True
        )
        self.main_window.classify_thread.start()
    
    def on_classify_progress(self, value, message):
        self.main_window.process_progress.setValue(value)
        self.main_window.process_log.append(message)
        self.main_window.status_bar.showMessage(message)
        
        self.main_window.process_status_label.setText(message)
    
    def on_classify_finished(self, output_file, stats):
        self.main_window.process_progress.setValue(100)
        
        self.main_window.process_status_label.setText(f"分类完成 - 匹配 {stats['matched']}/{stats['total']} 条")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #4CAF50;
            padding: 8px;
            background-color: #E8F5E9;
            border: 1px solid #4CAF50;
            border-radius: 4px;
        """)
        
        self.main_window.process_log.append(f"分类完成!")
        self.main_window.process_log.append(f"总记录数: {stats['total']}")
        self.main_window.process_log.append(f"匹配成功: {stats['matched']} 条")
        self.main_window.process_log.append(f"匹配种类: {stats['types']} 种")
        self.main_window.process_log.append(f"输出文件: {output_file}")
        
        mapping = stats['mapping']
        unmatched = stats.get('unmatched', {})
        
        total_rows = len(mapping)
        if unmatched:
            total_rows += 1
        
        self.main_window.process_table.setColumnCount(4)
        self.main_window.process_table.setHorizontalHeaderLabels(["标准岩性", "图片数", "分类数", "对应原始岩性"])
        self.main_window.process_table.setColumnWidth(0, 100)
        self.main_window.process_table.setColumnWidth(1, 80)
        self.main_window.process_table.setColumnWidth(2, 80)
        self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
        self.main_window.process_table.setRowCount(total_rows)
        
        row = 0
        for standard_rock, info in sorted(mapping.items(), key=lambda x: -x[1]['count']):
            lithologies_str = ", ".join(sorted(info['lithologies']))
            self.main_window.process_table.setItem(row, 0, QTableWidgetItem(standard_rock))
            self.main_window.process_table.setItem(row, 1, QTableWidgetItem(str(info['count'])))
            self.main_window.process_table.setItem(row, 2, QTableWidgetItem(str(len(info['lithologies']))))
            self.main_window.process_table.setItem(row, 3, QTableWidgetItem(lithologies_str))
            row += 1
        
        if unmatched:
            unmatched_str = ", ".join(sorted(unmatched.keys()))
            unmatched_count = sum(unmatched.values())
            self.main_window.process_table.setItem(row, 0, QTableWidgetItem("未发现分类"))
            self.main_window.process_table.setItem(row, 1, QTableWidgetItem(str(unmatched_count)))
            self.main_window.process_table.setItem(row, 2, QTableWidgetItem(str(len(unmatched))))
            self.main_window.process_table.setItem(row, 3, QTableWidgetItem(unmatched_str))
        
        self.main_window.process_table.resizeRowsToContents()
        
        self.main_window.process_btn.setEnabled(True)
        self.main_window.classify_btn.setEnabled(True)
        self.main_window.status_bar.showMessage("分类完成")
    
    def start_alteration_analysis(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择岩性分析后的JSON文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        app_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        project_root = os.path.dirname(app_dir)
        alteration_config_file = os.path.join(project_root, 'config', 'alteration_types.json')
        
        if not os.path.exists(alteration_config_file):
            QMessageBox.warning(self.main_window, "错误", f"找不到蚀变配置文件: {alteration_config_file}\n请检查config目录是否存在")
            return
        
        output_file = json_file.replace('.json', '_alteration.json')
        
        self.main_window.process_btn.setEnabled(False)
        self.main_window.classify_btn.setEnabled(False)
        if hasattr(self.main_window, 'alteration_btn'):
            self.main_window.alteration_btn.setEnabled(False)
        self.main_window.process_status_label.setText("正在初始化...")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #FF9800;
            padding: 8px;
            background-color: #FFF3E0;
            border: 1px solid #FF9800;
            border-radius: 4px;
        """)
        self.main_window.process_progress.setValue(0)
        self.main_window.process_log.clear()
        self.main_window.process_log.append("开始蚀变分析...")
        
        try:
            self.main_window.data_processor.progress_updated.disconnect()
            self.main_window.data_processor.processing_finished.disconnect()
            self.main_window.data_processor.error_occurred.disconnect()
        except TypeError:
            pass
        
        self.main_window.data_processor.progress_updated.connect(self.on_alteration_progress)
        self.main_window.data_processor.processing_finished.connect(self.on_alteration_finished)
        self.main_window.data_processor.error_occurred.connect(self.on_processing_error)
        
        from threading import Thread
        self.main_window.alteration_thread = Thread(
            target=self.main_window.data_processor.analyze_alteration,
            args=(json_file, alteration_config_file, output_file),
            daemon=True
        )
        self.main_window.alteration_thread.start()
    
    def on_alteration_progress(self, value, message):
        self.main_window.process_progress.setValue(value)
        self.main_window.process_log.append(message)
        self.main_window.status_bar.showMessage(message)
        self.main_window.process_status_label.setText(message)
    
    def on_alteration_finished(self, output_file, stats):
        import json
        
        self.main_window.process_progress.setValue(100)
        
        records_with_alt = stats.get('records_with_alteration', 0)
        total_records = stats.get('total_records', 0)
        total_alterations = stats.get('total_alterations', 0)
        alteration_types = stats.get('alteration_types', 0)
        
        self.main_window.process_status_label.setText(f"蚀变分析完成 - {records_with_alt}/{total_records} 条含蚀变")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #4CAF50;
            padding: 8px;
            background-color: #E8F5E9;
            border: 1px solid #4CAF50;
            border-radius: 4px;
        """)
        
        self.main_window.process_log.append(f"蚀变分析完成!")
        self.main_window.process_log.append(f"总记录数: {total_records}")
        self.main_window.process_log.append(f"含蚀变记录: {records_with_alt} 条")
        self.main_window.process_log.append(f"蚀变类型数: {alteration_types} 种")
        self.main_window.process_log.append(f"总蚀变次数: {total_alterations}")
        self.main_window.process_log.append(f"输出文件: {output_file}")
        
        with open(output_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        records = [(item.get('岩性名称', '') or '未分类', item.get('lithology', ''), item.get('蚀变类型', ''), item.get('lithology_description', '')) for item in data]
        records.sort(key=lambda x: (x[0], x[1]))
        
        self.main_window.process_table.setColumnCount(6)
        self.main_window.process_table.setHorizontalHeaderLabels(["标准岩性", "图片数", "分类数", "原始岩性", "蚀变类型", "岩性描述"])
        self.main_window.process_table.setColumnWidth(0, 100)
        self.main_window.process_table.setColumnWidth(1, 60)
        self.main_window.process_table.setColumnWidth(2, 60)
        self.main_window.process_table.setColumnWidth(3, 200)
        self.main_window.process_table.setColumnWidth(4, 120)
        self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
        
        rock_groups = {}
        for rock_name, lithology, alteration, description in records:
            if rock_name not in rock_groups:
                rock_groups[rock_name] = {}
            if lithology not in rock_groups[rock_name]:
                rock_groups[rock_name][lithology] = {'count': 0, 'alterations': set(), 'description': description}
            rock_groups[rock_name][lithology]['count'] += 1
            if alteration:
                rock_groups[rock_name][lithology]['alterations'].add(alteration)
        
        rows = []
        for rock_name in sorted(rock_groups.keys()):
            for lithology, info in rock_groups[rock_name].items():
                alterations_str = ", ".join(sorted(info['alterations'])) if info['alterations'] else "-"
                desc = info['description'] or "-"
                rows.append((rock_name, str(info['count']), "1", lithology, alterations_str, desc))
        
        self.main_window.process_table.setRowCount(len(rows))
        for row, (rock_name, count, _, lithology, alteration, description) in enumerate(rows):
            self.main_window.process_table.setItem(row, 0, QTableWidgetItem(rock_name))
            self.main_window.process_table.setItem(row, 1, QTableWidgetItem(count))
            self.main_window.process_table.setItem(row, 2, QTableWidgetItem("1"))
            self.main_window.process_table.setItem(row, 3, QTableWidgetItem(lithology))
            self.main_window.process_table.setItem(row, 4, QTableWidgetItem(alteration))
            self.main_window.process_table.setItem(row, 5, QTableWidgetItem(description))
        
        self.main_window.process_table.resizeRowsToContents()
        
        self.main_window.process_btn.setEnabled(True)
        self.main_window.classify_btn.setEnabled(True)
        if hasattr(self.main_window, 'alteration_btn'):
            self.main_window.alteration_btn.setEnabled(True)
        self.main_window.status_bar.showMessage("蚀变分析完成")
    
    def view_lithology_classification(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择岩性分类后的JSON文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        import json
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            QMessageBox.warning(self.main_window, "错误", f"无法读取文件: {str(e)}")
            return
        
        self.main_window.process_log.clear()
        self.main_window.process_log.append(f"已加载岩性分类结果: {json_file}")
        self.main_window.process_log.append(f"总记录数: {len(data)}")
        
        rock_groups = {}
        for item in data:
            rock_name = item.get('岩性名称', '') or '未分类'
            lithology = item.get('lithology', '')
            if rock_name not in rock_groups:
                rock_groups[rock_name] = {'lithologies': set(), 'count': 0}
            rock_groups[rock_name]['lithologies'].add(lithology)
            rock_groups[rock_name]['count'] += 1
        
        self.main_window.process_table.setColumnCount(4)
        self.main_window.process_table.setHorizontalHeaderLabels(["标准岩性", "图片数", "分类数", "对应原始岩性"])
        self.main_window.process_table.setColumnWidth(0, 100)
        self.main_window.process_table.setColumnWidth(1, 80)
        self.main_window.process_table.setColumnWidth(2, 80)
        self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
        
        sorted_rocks = sorted(rock_groups.items(), key=lambda x: -x[1]['count'])
        self.main_window.process_table.setRowCount(len(sorted_rocks))
        
        for row, (rock_name, info) in enumerate(sorted_rocks):
            lithologies_str = ", ".join(sorted(info['lithologies']))
            self.main_window.process_table.setItem(row, 0, QTableWidgetItem(rock_name))
            self.main_window.process_table.setItem(row, 1, QTableWidgetItem(str(info['count'])))
            self.main_window.process_table.setItem(row, 2, QTableWidgetItem(str(len(info['lithologies']))))
            self.main_window.process_table.setItem(row, 3, QTableWidgetItem(lithologies_str))
        
        self.main_window.process_table.resizeRowsToContents()
        
        self.main_window.process_status_label.setText(f"显示岩性分类 - 共 {len(rock_groups)} 种岩性")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #2196F3;
            padding: 8px;
            background-color: #E3F2FD;
            border: 1px solid #2196F3;
            border-radius: 4px;
        """)
        
        self.main_window.current_json_file = json_file
        self.main_window.status_bar.showMessage("已加载岩性分类结果")
    
    def view_alteration_analysis(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择蚀变分析后的JSON文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        import json
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            QMessageBox.warning(self.main_window, "错误", f"无法读取文件: {str(e)}")
            return
        
        self.main_window.process_log.clear()
        self.main_window.process_log.append(f"已加载蚀变分析结果: {json_file}")
        self.main_window.process_log.append(f"总记录数: {len(data)}")
        
        records = [(item.get('岩性名称', '') or '未分类', item.get('lithology', ''), item.get('蚀变类型', ''), item.get('lithology_description', '')) for item in data]
        records.sort(key=lambda x: (x[0], x[1]))
        
        self.main_window.process_table.setColumnCount(6)
        self.main_window.process_table.setHorizontalHeaderLabels(["标准岩性", "图片数", "分类数", "原始岩性", "蚀变类型", "岩性描述"])
        self.main_window.process_table.setColumnWidth(0, 100)
        self.main_window.process_table.setColumnWidth(1, 60)
        self.main_window.process_table.setColumnWidth(2, 60)
        self.main_window.process_table.setColumnWidth(3, 200)
        self.main_window.process_table.setColumnWidth(4, 120)
        self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
        
        rock_groups = {}
        for rock_name, lithology, alteration, description in records:
            if rock_name not in rock_groups:
                rock_groups[rock_name] = {}
            if lithology not in rock_groups[rock_name]:
                rock_groups[rock_name][lithology] = {'count': 0, 'alterations': set(), 'description': description}
            rock_groups[rock_name][lithology]['count'] += 1
            if alteration:
                rock_groups[rock_name][lithology]['alterations'].add(alteration)
        
        rows = []
        for rock_name in sorted(rock_groups.keys()):
            for lithology, info in rock_groups[rock_name].items():
                alterations_str = ", ".join(sorted(info['alterations'])) if info['alterations'] else "-"
                desc = info['description'] or "-"
                rows.append((rock_name, str(info['count']), "1", lithology, alterations_str, desc))
        
        self.main_window.process_table.setRowCount(len(rows))
        for row, (rock_name, count, _, lithology, alteration, description) in enumerate(rows):
            self.main_window.process_table.setItem(row, 0, QTableWidgetItem(rock_name))
            self.main_window.process_table.setItem(row, 1, QTableWidgetItem(count))
            self.main_window.process_table.setItem(row, 2, QTableWidgetItem("1"))
            self.main_window.process_table.setItem(row, 3, QTableWidgetItem(lithology))
            self.main_window.process_table.setItem(row, 4, QTableWidgetItem(alteration))
            self.main_window.process_table.setItem(row, 5, QTableWidgetItem(description))
        
        self.main_window.process_table.resizeRowsToContents()
        
        records_with_alt = sum(1 for item in data if item.get('蚀变类型'))
        self.main_window.process_status_label.setText(f"显示蚀变分析 - {records_with_alt}/{len(data)} 条含蚀变")
        self.main_window.process_status_label.setStyleSheet("""
            font-size: 14px;
            font-weight: bold;
            color: #9C27B0;
            padding: 8px;
            background-color: #F3E5F5;
            border: 1px solid #9C27B0;
            border-radius: 4px;
        """)
        
        self.main_window.current_json_file = json_file
        self.main_window.status_bar.showMessage("已加载蚀变分析结果")
    
    def export_to_excel(self):
        from PyQt6.QtWidgets import QFileDialog
        table = self.main_window.process_table
        
        if table.rowCount() == 0:
            QMessageBox.warning(self.main_window, "警告", "表格中没有数据可导出")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self.main_window, "导出Excel文件", "", "Excel文件 (*.xlsx)"
        )
        if not file_path:
            return
        
        if not file_path.endswith('.xlsx'):
            file_path += '.xlsx'
        
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据导出"
            
            header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for col in range(table.columnCount()):
                header_item = table.horizontalHeaderItem(col)
                header_text = header_item.text() if header_item else ""
                cell = ws.cell(row=1, column=col+1, value=header_text)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            for row in range(table.rowCount()):
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    cell_value = item.text() if item else ""
                    ws.cell(row=row+2, column=col+1, value=cell_value)
            
            for col in range(table.columnCount()):
                max_length = 0
                col_letter = openpyxl.utils.get_column_letter(col+1)
                column = ws.column_dimensions[col_letter]
                for cell in ws[col_letter]:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                column.width = min(max_length + 2, 60)
            
            wb.save(file_path)
            QMessageBox.information(self.main_window, "导出成功", f"数据已导出到:\n{file_path}")
            self.main_window.process_log.append(f"已导出数据到: {file_path}")
        except Exception as e:
            QMessageBox.warning(self.main_window, "错误", f"导出失败: {str(e)}")