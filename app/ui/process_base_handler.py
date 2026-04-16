import os
import json
from PyQt6.QtWidgets import QMessageBox, QTableWidgetItem


class ProcessBaseHandler:
    def __init__(self, main_window):
        self.main_window = main_window
    
    def select_source_directory(self):
        from PyQt6.QtWidgets import QFileDialog
        folder = QFileDialog.getExistingDirectory(self.main_window, "选择数据目录")
        if folder:
            self.main_window.source_directory = folder
            self.main_window.source_path_label.setText(os.path.basename(folder))
            self.main_window.process_log.append(f"已选择数据源: {folder}")
    
    def select_output_directory(self):
        from PyQt6.QtWidgets import QFileDialog
        folder = QFileDialog.getExistingDirectory(self.main_window, "选择输出目录")
        if folder:
            self.main_window.output_directory = folder
            self.main_window.output_path_label.setText(os.path.basename(folder))
            self.main_window.process_log.append(f"已选择输出目录: {folder}")
    
    def show_process_stats(self):
        from PyQt6.QtWidgets import QFileDialog
        json_file, _ = QFileDialog.getOpenFileName(
            self.main_window, "选择处理结果文件", "", "JSON文件 (*.json)"
        )
        if not json_file:
            return
        
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            QMessageBox.warning(self.main_window, "错误", f"无法读取文件: {str(e)}")
            return
        
        self.main_window.process_log.clear()
        self.main_window.process_log.append(f"正在加载: {json_file}")
        
        lithology_file = json_file.replace('image_descriptions.json', 'lithology_descriptions.json')
        lithology_data = []
        if os.path.exists(lithology_file):
            try:
                with open(lithology_file, 'r', encoding='utf-8') as f:
                    lithology_data = json.load(f)
            except:
                pass
        
        if lithology_data:
            lithology_data_sorted = sorted(lithology_data, key=lambda x: x.get('id', 0))
            
            image_counts = {}
            for item in data:
                desc_id = item.get('lithology_description_id')
                if desc_id is not None:
                    image_counts[desc_id] = image_counts.get(desc_id, 0) + 1
            
            self.main_window.process_table.setColumnCount(5)
            self.main_window.process_table.setHorizontalHeaderLabels(["项目", "钻孔", "岩性名称", "深度范围", "图片数"])
            self.main_window.process_table.setColumnWidth(0, 120)
            self.main_window.process_table.setColumnWidth(1, 100)
            self.main_window.process_table.setColumnWidth(2, 150)
            self.main_window.process_table.setColumnWidth(3, 100)
            self.main_window.process_table.setColumnWidth(4, 80)
            self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
            self.main_window.process_table.setRowCount(len(lithology_data_sorted))
            
            for i, lith in enumerate(lithology_data_sorted):
                proj = lith.get('project', '')
                borehole = lith.get('borehole', '')
                lith_name = lith.get('lithology', '')
                start = lith.get('start_depth', 0)
                end = lith.get('end_depth', 0)
                depth_range = f"{start}-{end}"
                count = image_counts.get(lith.get('id', i), 0)
                
                self.main_window.process_table.setItem(i, 0, QTableWidgetItem(proj))
                self.main_window.process_table.setItem(i, 1, QTableWidgetItem(borehole))
                self.main_window.process_table.setItem(i, 2, QTableWidgetItem(lith_name))
                self.main_window.process_table.setItem(i, 3, QTableWidgetItem(depth_range))
                self.main_window.process_table.setItem(i, 4, QTableWidgetItem(str(count)))
            
            self.main_window.process_table.resizeRowsToContents()
            total_images = sum(image_counts.values())
            self.main_window.process_log.append(f"共 {len(lithology_data)} 条岩性记录, {total_images} 张图片")
            self.main_window.status_bar.showMessage("统计完成")
        else:
            project_stats = {}
            for item in data:
                proj = item.get('project', '未知')
                if proj not in project_stats:
                    project_stats[proj] = {'count': 0, 'lithologies': set()}
                project_stats[proj]['count'] += 1
                lith = item.get('lithology', '')
                if lith:
                    project_stats[proj]['lithologies'].add(lith)
            
            self.main_window.process_table.setColumnCount(3)
            self.main_window.process_table.setHorizontalHeaderLabels(["项目", "图片数", "岩性种类"])
            self.main_window.process_table.setColumnWidth(0, 150)
            self.main_window.process_table.setColumnWidth(1, 80)
            self.main_window.process_table.horizontalHeader().setStretchLastSection(True)
            self.main_window.process_table.setRowCount(len(project_stats))
            
            for i, (proj, stats) in enumerate(sorted(project_stats.items())):
                self.main_window.process_table.setItem(i, 0, QTableWidgetItem(proj))
                self.main_window.process_table.setItem(i, 1, QTableWidgetItem(str(stats['count'])))
                self.main_window.process_table.setItem(i, 2, QTableWidgetItem(str(len(stats['lithologies']))))
            
            self.main_window.process_table.resizeRowsToContents()
            self.main_window.process_log.append(f"共 {len(project_stats)} 个项目")
            self.main_window.status_bar.showMessage("统计完成")
    
    def export_to_excel(self):
        from PyQt6.QtWidgets import QFileDialog
        from PyQt6.QtWidgets import QProgressDialog
        from PyQt6.QtCore import Qt
        
        if self.main_window.process_table.rowCount() == 0:
            QMessageBox.warning(self.main_window, "警告", "没有可导出的数据")
            return
        
        output_file, _ = QFileDialog.getSaveFileName(
            self.main_window, "导出Excel", "", "Excel文件 (*.xlsx)"
        )
        if not output_file:
            return
        
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            
            wb = Workbook()
            ws = wb.active
            ws.title = "处理结果"
            
            headers = []
            for col in range(self.main_window.process_table.columnCount()):
                item = self.main_window.process_table.horizontalHeaderItem(col)
                headers.append(item.text() if item else "")
            
            header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for col, header in enumerate(headers):
                cell = ws.cell(row=1, column=col + 1, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            
            for row in range(self.main_window.process_table.rowCount()):
                for col in range(self.main_window.process_table.columnCount()):
                    item = self.main_window.process_table.item(row, col)
                    value = item.text() if item else ""
                    cell = ws.cell(row=row + 2, column=col + 1, value=value)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border
            
            for col in range(len(headers)):
                ws.column_dimensions[chr(65 + col)].width = 20
            
            wb.save(output_file)
            
            self.main_window.process_log.append(f"已导出到: {output_file}")
            QMessageBox.information(self.main_window, "导出成功", f"数据已导出到:\n{output_file}")
            
        except Exception as e:
            QMessageBox.warning(self.main_window, "导出失败", f"导出时出错:\n{str(e)}")
